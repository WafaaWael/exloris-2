"""Render the project flowcharts and UML diagrams as dark-theme PNG figures."""

from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent / "rendered"
DOCX = ROOT / "Documentation" / "Exloris-Technical-Documentation.docx"
OUT.mkdir(parents=True, exist_ok=True)

W, H = 1600, 1050
BG = "#0D1117"
NODE = "#EEEAFE"
NODE_ALT = "#E3DCFF"
INK = "#282439"
EDGE = "#586174"
ACCENT = "#B89CFF"
MUTED = "#AAB3C2"
GREEN = "#DDF5E7"
RED = "#FCE1E1"

FONT_CANDIDATES = [
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]
BOLD_CANDIDATES = [
    "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    "/Library/Fonts/Arial Bold.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
]
NODE_RECTS = []


def font(size, bold=False):
    candidates = BOLD_CANDIDATES if bold else FONT_CANDIDATES
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def canvas(width=W, height=H):
    NODE_RECTS.clear()
    im = Image.new("RGB", (width, height), BG)
    return im, ImageDraw.Draw(im)


def wrap_lines(draw, text, fnt, max_width):
    words = text.split()
    lines, line = [], ""
    for word in words:
        attempt = (line + " " + word).strip()
        if line and draw.textbbox((0, 0), attempt, font=fnt)[2] > max_width:
            lines.append(line)
            line = word
        else:
            line = attempt
    if line:
        lines.append(line)
    return lines


def node(draw, cx, cy, text, width=300, height=94, shape="box", fill=NODE, size=36):
    x0, y0, x1, y1 = cx-width//2, cy-height//2, cx+width//2, cy+height//2
    if shape == "diamond":
        pts = [(cx, y0), (x1, cy), (cx, y1), (x0, cy)]
        draw.polygon(pts, fill=fill, outline=ACCENT, width=4)
    elif shape == "pill":
        draw.rounded_rectangle((x0,y0,x1,y1), radius=height//2, fill=fill, outline=ACCENT, width=4)
    else:
        draw.rectangle((x0,y0,x1,y1), fill=fill, outline=ACCENT, width=4)
    NODE_RECTS.append((x0,y0,x1,y1,shape,cx,cy))
    fnt = font(size)
    lines = []
    for part in text.split("\n"):
        lines.extend(wrap_lines(draw, part, fnt, width-32))
    line_h = size + 8
    height = max(height, len(lines)*line_h+28)
    y0, y1 = cy-height//2, cy+height//2
    if shape == "diamond":
        pts = [(cx, y0), (x1, cy), (cx, y1), (x0, cy)]
        draw.polygon(pts, fill=fill, outline=ACCENT, width=4)
    elif shape == "pill":
        draw.rounded_rectangle((x0,y0,x1,y1), radius=height//2, fill=fill, outline=ACCENT, width=4)
    else:
        draw.rectangle((x0,y0,x1,y1), fill=fill, outline=ACCENT, width=4)
    NODE_RECTS[-1]=(x0,y0,x1,y1,shape,cx,cy)
    y = cy - (len(lines)*line_h)//2
    for line in lines:
        bbox = draw.textbbox((0,0), line, font=fnt)
        draw.text((cx-(bbox[2]-bbox[0])/2, y), line, font=fnt, fill=INK)
        y += line_h
    return (x0,y0,x1,y1)


def arrow(draw, points, label=None, label_at=None, color=EDGE, dashed=False):
    if len(points) < 2:
        return
    points = list(points)

    def trim(point, toward):
        px,py=point
        for x0,y0,x1,y1,shape,cx,cy in reversed(NODE_RECTS):
            if x0 < px < x1 and y0 < py < y1:
                dx,dy=toward[0]-cx,toward[1]-cy
                if not dx and not dy: return point
                if shape=="diamond":
                    scale=1.0/max(abs(dx)/max(1,(x1-x0)/2)+abs(dy)/max(1,(y1-y0)/2),1e-6)
                else:
                    scale=min((x1-x0)/2/abs(dx) if dx else float("inf"),(y1-y0)/2/abs(dy) if dy else float("inf"))
                return (int(cx+dx*scale),int(cy+dy*scale))
        return point

    points[0]=trim(points[0],points[1])
    points[-1]=trim(points[-1],points[-2])
    for (x1,y1),(x2,y2) in zip(points, points[1:]):
        if dashed and x1 == x2:
            yy = y1
            step = 18 if y2 >= y1 else -18
            while (yy < y2 if step > 0 else yy > y2):
                nxt = yy + step
                if (step > 0 and nxt > y2) or (step < 0 and nxt < y2): nxt = y2
                draw.line((x1,yy,x2,nxt), fill=color, width=4)
                yy = nxt + step
        elif dashed and y1 == y2:
            xx = x1
            step = 18 if x2 >= x1 else -18
            while (xx < x2 if step > 0 else xx > x2):
                nxt=xx+step
                if (step > 0 and nxt > x2) or (step < 0 and nxt < x2): nxt=x2
                draw.line((xx,y1,nxt,y2), fill=color, width=4)
                xx=nxt+step
        else:
            draw.line((x1,y1,x2,y2), fill=color, width=4, joint="curve")
    (x1,y1),(x2,y2)=points[-2],points[-1]
    import math
    a=math.atan2(y2-y1,x2-x1); L=18; spread=.55
    p1=(x2-L*math.cos(a-spread),y2-L*math.sin(a-spread))
    p2=(x2-L*math.cos(a+spread),y2-L*math.sin(a+spread))
    draw.polygon([(x2,y2),p1,p2],fill=color)
    if label:
        lx,ly=label_at or ((x1+x2)//2,(y1+y2)//2)
        fnt=font(23,True)
        max_width=max(125,min(430,abs(x2-x1)-24 if abs(x2-x1)>145 else 260))
        lines=wrap_lines(draw,label,fnt,max_width)
        boxes=[draw.textbbox((0,0),line,font=fnt) for line in lines]
        tw=max(b[2]-b[0] for b in boxes); line_h=31; th=len(lines)*line_h
        draw.rounded_rectangle((lx-tw//2-10,ly-th//2-7,lx+tw//2+10,ly+th//2+7),radius=7,fill=BG)
        yy=ly-th//2-2
        for line,b in zip(lines,boxes):
            draw.text((lx-(b[2]-b[0])/2,yy),line,font=fnt,fill=MUTED); yy+=line_h


def save(im, name):
    path=OUT/(name+".png")
    im.save(path,optimize=True)
    return path


def flow_overview():
    im,d=canvas(1800,1120)
    node(d,900,90,"Scenes and prefabs",390,82)
    node(d,260,270,"Menu / player / map",350,86)
    node(d,900,270,"BoatLevel",310,86)
    node(d,1500,270,"Nilometer",310,86)
    node(d,180,480,"SceneManager /\nserialized UI callbacks",300,108)
    node(d,540,480,"Avatar selection /\nSaveManager",330,108)
    node(d,900,480,"BoatLevelFlow",330,88)
    node(d,1500,480,"NilometerUI",310,88)
    node(d,850,690,"BoatInput / BoatController /\nCoin / Obstacle",420,112)
    node(d,1300,690,"Typed EventBus",300,88)
    node(d,1600,690,"WaterOperator /\nWaterAnimator",280,108)
    node(d,540,850,"SaveSystem",300,86)
    node(d,1070,970,"AudioService",290,86)
    node(d,1350,970,"BoatLevelVfx",250,86)
    node(d,1640,900,"Shared DialoguePanel",300,86)
    for a,b in [((900,131),(260,227)),((900,131),(900,227)),((900,131),(1500,227)),((260,313),(180,426)),((260,313),(540,426)),((900,313),(900,436)),((1500,313),(1500,436)),((900,524),(850,634)),((900,524),(1300,646)),((1500,524),(1600,636)),((540,534),(540,807)),((1300,734),(1070,927)),((1300,734),(1380,927))]:
        arrow(d,[a,b])
    arrow(d,[(1500,480),(1780,480),(1780,830),(1640,857)])
    arrow(d,[(1650,900),(1760,900),(1760,760),(1300,760),(1300,734)])
    return save(im,"01-project-overview")


def flow_startup():
    im,d=canvas()
    node(d,800,80,"Player launches build",360,76,shape="pill")
    node(d,800,190,"Unity selects first enabled\nbuild scene",430,100)
    node(d,800,325,"Main menu scene",330,82)
    node(d,800,455,"Awake / OnEnable",330,82)
    node(d,800,585,"Start callbacks",330,82)
    node(d,475,740,"Save / scene / sound /\naudio services initialize",430,105)
    node(d,1125,740,"SceneOrgniser invokes\nserialized start event",430,105)
    node(d,800,900,"Menu interaction ready → player selection",620,86)
    for y1,y2 in [(118,140),(240,284),(366,414),(496,544),(626,684)]: arrow(d,[(800,y1),(800,y2)])
    arrow(d,[(800,626),(475,626),(475,686)])
    arrow(d,[(800,626),(1125,626),(1125,686)])
    arrow(d,[(475,793),(475,850),(800,850),(800,856)])
    arrow(d,[(1125,793),(1125,850),(800,850),(800,856)])
    return save(im,"02-startup-flow")


def flow_gameplay():
    im,d=canvas(1800,1200)
    node(d,190,110,"Intro dialogue",260,78)
    node(d,520,110,"Tutorial",250,78)
    node(d,870,110,"Playing\nboat and input enabled",330,100)
    node(d,1330,110,"Keyboard / pointer input",360,86)
    node(d,1330,320,"Emit lane step",290,80)
    node(d,990,320,"BoatController moves\nwithin three lanes",370,100)
    node(d,570,320,"Trigger contact?",300,138,"diamond")
    node(d,190,540,"Coin collected once\ncount increments",350,106)
    node(d,610,540,"No contact\nprogress and speed update",380,106)
    node(d,1050,540,"Obstacle response\npushback / block / bounce",400,106)
    node(d,190,750,"Coin event and SFX",320,84)
    node(d,610,750,"HUD / VFX / audio feedback",380,88)
    node(d,1050,750,"Eligible hit updates stars\nand publishes hit event",400,108)
    node(d,610,930,"Near dock?",270,136,"diamond")
    node(d,1210,930,"Continue boat loop",330,84)
    node(d,610,1100,"Arrival flow\nDisable input • autopilot • dialogue • completion",520,112)
    node(d,1280,1100,"Replay reloads scene\nNext currently does not load",420,108)
    for a,b in [((320,110),(395,110)),((645,110),(705,110)),((1035,110),(1150,110)),((1330,153),(1330,280)),((1150,320),(1175,320)),((805,320),(720,320)),((570,389),(570,487)),((570,340),(365,500)),((570,340),(850,500)),((190,593),(190,707)),((1050,593),(1050,696)),((365,750),(420,750)),((610,804),(610,862)),((745,930),(1045,930)),((610,998),(610,1044)),((870,1100),(1070,1100))]: arrow(d,[a,b])
    arrow(d,[(420,540),(390,540),(390,930),(475,930)])
    arrow(d,[(850,750),(800,750)])
    arrow(d,[(1375,930),(1580,930),(1580,153),(1330,153)],"No",(1480,900))
    arrow(d,[(1280,1154),(1700,1154),(1700,153),(1330,153)])
    return save(im,"03-boat-gameplay-flow")


def flow_transitions():
    im,d=canvas(1800,1100)
    for x,t in [(180,"Main menu"),(500,"Player"),(820,"Map"),(1140,"BoatLevel"),(1470,"Nilometer")]: node(d,x,180,t,250,82)
    for x1,x2 in [(305,375),(625,695),(945,1015),(1265,1345)]: arrow(d,[(x1,180),(x2,180)])
    node(d,1150,390,"Replay\nreload active scene",300,100)
    node(d,1470,390,"Completion Next button",320,86)
    node(d,1470,570,"Current handler does not\ninvoke LoadNextSection",380,100,fill=RED)
    node(d,490,770,"BoatLevelFlow.LoadNextSection",400,86)
    node(d,900,770,"Load configured scene name",380,86)
    node(d,1300,770,"Enabled destination?",300,128,"diamond")
    node(d,1100,980,"Unavailable",280,78,fill=RED)
    node(d,1510,980,"Load destination",280,78,fill=GREEN)
    arrow(d,[(1140,222),(1140,325)])
    arrow(d,[(1470,222),(1470,347)])
    arrow(d,[(1470,433),(1470,520)])
    arrow(d,[(690,770),(710,770)])
    arrow(d,[(1090,770),(1150,770)])
    arrow(d,[(1300,834),(1300,920),(1100,920),(1100,940)],"No",(1200,906))
    arrow(d,[(1300,834),(1300,920),(1510,920),(1510,940)],"Yes",(1410,906))
    return save(im,"04-scene-transitions")


def flow_events():
    im,d=canvas(1800,1100)
    sources=[(260,220,"BoatLevelFlow"),(260,430,"BoatController"),(260,640,"NilometerUI"),(260,850,"BoatLevelUI / OptionsPanel")]
    ev=[(760,320,"State / progress /\ncoin / hit events"),(760,690,"SFX events")]
    consumers=[(1250,500,"Typed EventBus"),(1570,330,"UI / Audio / VFX"),(1570,650,"VoiceActivityEvent"),(1570,850,"CharacterExpressions")]
    for x,y,t in sources: node(d,x,y,t,390,82)
    for x,y,t in ev: node(d,x,y,t,430,100)
    for x,y,t in consumers: node(d,x,y,t,400,86)
    arrow(d,[(455,220),(545,320)])
    arrow(d,[(455,430),(545,320)])
    arrow(d,[(455,220),(500,220),(500,690),(545,690)])
    arrow(d,[(455,640),(545,690)])
    arrow(d,[(455,850),(500,850),(500,690),(545,690)])
    arrow(d,[(975,320),(1050,320),(1050,500)])
    arrow(d,[(975,690),(1050,690),(1050,520)])
    arrow(d,[(1450,500),(1370,330)])
    arrow(d,[(1570,373),(1570,607)])
    arrow(d,[(1570,693),(1570,807)])
    return save(im,"05-typed-event-flow")


def flow_completion():
    im,d=canvas(1700,1100)
    node(d,450,110,"Boat reaches dock approach",430,82)
    node(d,450,245,"Autopilot starts\ninput disabled",360,100)
    node(d,450,395,"Boat reaches stop point",380,82)
    node(d,450,535,"Arrival delay and dialogue",380,82)
    node(d,450,680,"Dialogue finished callback\nCompleteArrival and win SFX",430,105)
    node(d,450,835,"Completion card\nstars and coins",360,98)
    node(d,220,1010,"Replay reloads active scene",370,82)
    node(d,700,1010,"Next button currently\ndoes not load another scene",390,96,fill=RED)
    node(d,1260,330,"Stars reach zero",340,82)
    node(d,1260,510,"No explicit boat failure\nor game-over state",390,104,fill=RED)
    for y1,y2 in [(151,195),(295,354),(436,494),(577,628),(733,786),(884,968)]: arrow(d,[(450,y1),(450,y2)])
    arrow(d,[(1260,372),(1260,458)])
    return save(im,"06-victory-failure-restart")


def class_box(draw,cx,cy,w,title,items,fill=NODE):
    f2=font(30)
    wrapped=[]
    for item in items:
        wrapped.extend(wrap_lines(draw,item,f2,w-36))
    h=88+len(wrapped)*42; x0=cx-w//2; y0=cy-h//2; x1=cx+w//2; y1=cy+h//2
    draw.rectangle((x0,y0,x1,y1),fill=fill,outline=ACCENT,width=4)
    draw.rectangle((x0,y0,x1,y0+54),fill=NODE_ALT,outline=ACCENT,width=3)
    f=font(34,True); bb=draw.textbbox((0,0),title,font=f); draw.text((cx-(bb[2]-bb[0])/2,y0+10),title,font=f,fill=INK)
    y=y0+70
    for it in wrapped:
        draw.text((x0+18,y),it,font=f2,fill=INK); y+=42
    NODE_RECTS.append((x0,y0,x1,y1,"box",cx,cy))
    return (x0,y0,x1,y1)


def class_boat():
    im,d=canvas(1800,1250)
    class_box(d,900,190,500,"BoatLevelFlow",["State • Stars • Coins","Start / Replay / Complete","Owns level state and score"])
    class_box(d,360,450,390,"BoatInput",["IsInputEnabled","OnLaneStep event"])
    class_box(d,900,480,430,"BoatController",["Lane • IsMoving • IsBlocked","StepLane / SetProgress","Autopilot / collision callbacks"])
    class_box(d,1440,450,400,"Coin / Obstacle",["Trigger contacts","Collect once / obstacle response"])
    class_box(d,360,820,420,"EventPayloads",["Level / progress / coin","hit / blocked / SFX events"])
    class_box(d,900,820,360,"EventBus",["Typed publish / subscribe"])
    class_box(d,1450,820,430,"UI / Audio / VFX",["BoatLevelUI","AudioService / BoatLevelVfx"])
    arrow(d,[(360,450),(900,480)])
    arrow(d,[(900,288),(900,382)])
    arrow(d,[(1115,450),(1240,450)])
    arrow(d,[(900,578),(900,722)])
    arrow(d,[(570,820),(720,820)])
    arrow(d,[(1080,820),(1235,820)])
    return save(im,"07-boat-class-diagram")


def class_nilometer():
    im,d=canvas(1900,1280)
    class_box(d,900,190,530,"NilometerUI",["SubmitAnswer / Hint","Quiz flow and feedback","Publishes SFX events"])
    class_box(d,900,485,470,"WaterOperator",["Level count / current level","Answer validation / stage rules"])
    class_box(d,350,485,420,"Tail answer panels",["Serialized button bindings","Answer selection"])
    class_box(d,1500,485,430,"SharedDialoguePanel",["PlayLines / Hide","Finished and voice callbacks"])
    class_box(d,900,785,420,"WaterAnimator",["Animate water height","Completion callbacks"])
    class_box(d,350,850,430,"CharacterExpressions",["Mood and reaction"])
    class_box(d,1500,850,360,"StarsDisplay",["Star feedback"])
    class_box(d,900,1050,300,"EventBus",["SFX events"])
    arrow(d,[(900,190),(350,485)])
    arrow(d,[(900,190),(900,485)])
    arrow(d,[(900,485),(900,785)])
    arrow(d,[(900,190),(1360,190),(1360,378),(1500,378)])
    arrow(d,[(900,190),(110,190),(110,850),(135,850)])
    arrow(d,[(900,190),(1790,190),(1790,780),(1500,780)])
    arrow(d,[(900,190),(1840,190),(1840,950),(900,950),(900,985)])
    return save(im,"08-nilometer-class-diagram")


def class_persistence():
    im,d=canvas(1800,900)
    class_box(d,300,360,360,"player",["Level / Step / Gender","Active scene player"])
    class_box(d,730,360,390,"SaveManager",["CurrentData","Register / save / load","Player lifecycle"])
    class_box(d,1170,360,390,"SaveSystem",["JSON serialization","AES + HMAC envelope","File read / write"])
    class_box(d,1580,360,320,"PlayerData",["Level / Step / Gender"])
    class_box(d,1170,700,360,"SaveSecrets",["Key provider","Secret value omitted"],fill=RED)
    arrow(d,[(300,360),(730,360)])
    arrow(d,[(730,360),(1170,360)])
    arrow(d,[(1170,360),(1580,360)])
    arrow(d,[(1170,490),(1170,640)])
    return save(im,"09-persistence-class-diagram")


def state_boat():
    im,d=canvas(1800,1050)
    positions={"Intro":(150,500),"Tutorial":(440,500),"Playing":(740,500),"Paused":(740,180),"Settings":(740,820),"Arriving":(1030,500),"Arrived":(1320,500),"Complete":(1620,500)}
    for t,(x,y) in positions.items(): node(d,x,y,t,250,76,shape="pill" if t in ("Intro","Complete") else "box")
    arrow(d,[(275,500),(315,500)],"CompleteIntro",(295,425))
    arrow(d,[(565,500),(615,500)],"StartPlaying",(590,425))
    arrow(d,[(865,500),(905,500)],"dock threshold",(885,425))
    arrow(d,[(1155,500),(1195,500)],"autopilot stop + delay",(1175,425))
    arrow(d,[(1445,500),(1495,500)],"dialogue complete",(1470,425))
    arrow(d,[(740,458),(740,222)],"OpenHelp",(840,340))
    arrow(d,[(740,222),(740,458)],"StartPlaying",(635,340))
    arrow(d,[(740,542),(740,778)],"OpenSettings",(845,660))
    arrow(d,[(740,778),(740,542)],"CloseSettings",(630,660))
    arrow(d,[(1620,542),(1620,960),(150,960),(150,542)],"Replay reloads scene",(930,920))
    return save(im,"10-boat-state-diagram")


def sequence(name, actors, messages, width=1800, height=1250, branches=()):
    im,d=canvas(width,height)
    xs=[int(100+(width-200)*i/(len(actors)-1)) for i in range(len(actors))]
    top=90; head_w=min(215,int(width/(len(actors)+0.5)))
    for x,label in zip(xs,actors):
        node(d,x,top,label,head_w,84,size=30)
        y=top+43
        while y<height-70:
            d.line((x,y,x,y+18),fill=EDGE,width=2)
            y+=34
    y0=190; gap=76
    for item in messages:
        fr,to,label,kind=item[:4]
        y=y0
        x1,x2=xs[fr],xs[to]
        col=ACCENT if kind=="reply" else EDGE
        if kind=="self":
            d.line((x1,y,x1+70,y),fill=col,width=3)
            d.line((x1+70,y,x1+70,y+34),fill=col,width=3)
            d.line((x1+70,y+34,x1,y+34),fill=col,width=3)
            arrow(d,[(x1+1,y+34),(x1,y+34)],color=col)
            d.text((x1+75,y-22),label,font=font(29),fill=NODE)
        else:
            arrow(d,[(x1,y),(x2,y)],label,(int((x1+x2)/2),y-30),color=col,dashed=kind=="reply")
        y0+=gap
    return save(im,name)


def sequence_boat():
    actors=["User","BoatInput","BoatController","BoatLevelFlow","EventBus","BoatLevelUI","BoatLevelVfx","AudioService"]
    msgs=[(0,1,"keyboard / swipe / tap","call"),(1,3,"OnLaneStep(direction)","call"),(3,2,"StepLane(direction)","call"),(2,3,"OnObstacleHit(first hit)","call"),(3,4,"PlaySfxEvent(Bump)","call"),(3,3,"decrement Stars if eligible","self"),(3,4,"BoatHitObstacleEvent","call"),(4,5,"update stars","reply"),(4,6,"impact / shake","reply"),(4,7,"play registered SFX","reply")]
    return sequence("11-sequence-boat-obstacle",actors,msgs,height=1050)


def sequence_nilometer():
    actors=["User","Question button","NilometerUI","WaterOperator","DialoguePanel","EventBus"]
    msgs=[(0,1,"select answer","call"),(1,2,"SubmitAnswer(level, answer)","call"),(2,3,"IsCorrectAnswer","call"),(2,5,"Correct / Wrong SFX","call"),(3,2,"correct result","reply"),(2,2,"feedback and continue","self"),(2,4,"first miss: play almost line","call"),(4,2,"finished; resume question","reply"),(2,4,"second miss: play sorry line","call"),(4,2,"finished; advance level","reply")]
    return sequence("12-sequence-nilometer-answer",actors,msgs,height=1050)


def sequence_save():
    actors=["Player","SaveManager","SaveSystem","player.save"]
    msgs=[(0,1,"RegisterPlayer / SaveCurrentPlayer","call"),(1,2,"SavePlayer(player)","call"),(2,2,"serialize JSON; AES encrypt","self"),(2,2,"HMAC envelope; Base64","self"),(2,3,"write file","call"),(1,2,"LoadPlayer()","call"),(2,3,"read file","call"),(2,2,"verify header and HMAC","self"),(2,2,"decrypt and deserialize","self"),(2,1,"return PlayerData or null","reply")]
    return sequence("13-sequence-save-load",actors,msgs,height=1050)


def build_images():
    paths=[flow_overview(),flow_startup(),flow_gameplay(),flow_transitions(),flow_events(),flow_completion(),class_boat(),class_nilometer(),class_persistence(),state_boat(),sequence_boat(),sequence_nilometer(),sequence_save()]
    return paths


def embed_images(paths):
    doc=Document(str(DOCX))
    names=["A. Project overview","B. Startup and initialization","C. Main gameplay loop — boat level","D. Scene transitions","E. Typed gameplay event flow","F. Victory, failure, and restart","1. Boat system class diagram","2. Nilometer class diagram","3. Persistence class diagram","4. Boat state diagram","5. Sequence diagram — boat obstacle hit","6. Sequence diagram — Nilometer answer","7. Sequence diagram — save/load"]
    paragraphs=doc.paragraphs
    targets=[]
    current=None
    for i,p in enumerate(paragraphs):
        if p.style.name=="Heading 1" and p.text in ("Editable flowcharts","UML-style diagrams"):
            current=p.text
        elif p.style.name=="Heading 1" and p.text not in ("Editable flowcharts","UML-style diagrams"):
            current=None
        if p.style.name=="Heading 2" and p.text in names and current:
            targets.append((names.index(p.text),i,p.text))
    if len(targets)!=len(paths):
        raise RuntimeError(f"Expected {len(paths)} diagram headings, found {len(targets)}")
    remove=[]
    for slot,(_,start,title) in enumerate(targets):
        # Give each diagram a full-width page so Word does not shrink adjacent
        # figures until their labels become difficult to read.
        paragraphs[start].paragraph_format.page_break_before=True
        end=len(paragraphs)
        for k in range(start+1,len(paragraphs)):
            if paragraphs[k].style.name in ("Heading 1","Heading 2"):
                end=k; break
        block=[p for p in paragraphs[start+1:end] if p.style.name=="Code Block"]
        if not block:
            picture=None
            for shape in doc.inline_shapes:
                descr=shape._inline.docPr.get("descr") or ""
                if descr==f"Visual diagram: {title}":
                    picture=shape
                    break
            if picture is None:
                raise RuntimeError(f"No Mermaid source or visual diagram found under {title}")
            blip=picture._inline.graphic.graphicData.pic.blipFill.blip
            rid=blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            doc.part.related_parts[rid]._blob=paths[slot].read_bytes()
            continue
        marker=block[0]
        image_p=marker.insert_paragraph_before()
        image_p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        image_p.paragraph_format.keep_together=True
        run=image_p.add_run()
        shape=run.add_picture(str(paths[slot]),width=Inches(6.15))
        shape._inline.docPr.set("descr",f"Visual diagram: {title}")
        remove.extend(block)
    for p in remove:
        p._element.getparent().remove(p._element)
    doc.save(str(DOCX))


if __name__=="__main__":
    images=build_images()
    embed_images(images)
    print(f"Rendered and embedded {len(images)} visual diagrams in {DOCX}")
